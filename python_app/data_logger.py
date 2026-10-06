"""
data_logger.py — Modul penyimpanan data ke file Excel (.xlsx) dengan dukungan Offline Buffering & Sync Status.

Modul ini menangani:
  - Pembuatan file Excel (.xlsx) otomatis dengan format rapi (header berwarna, auto-width, freeze pane)
  - Penambahan data baru ke file Excel secara berkala beserta kolom Status_Sync (SYNCED / PENDING)
  - Pengambilan data berstatus PENDING saat internet offline untuk disinkronkan kembali
  - Pembaruan status baris menjadi SYNCED setelah berhasil terkirim ke database
  - Validasi header kolom agar nilai selalu masuk ke kolom yang tepat
  - Penanganan file lock jika Excel sedang dibuka user
"""

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config


def get_excel_path():
    """Mendapatkan path lengkap file Excel (.xlsx)."""
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), config.EXCEL_FILENAME)


# Alias untuk kompatibilitas
def get_csv_path():
    return get_excel_path()


def get_headers():
    """Mendapatkan daftar header kolom sesuai konfigurasi REGISTER_MAP terkini dan Status_Sync."""
    headers = ["Timestamp", "Tanggal", "Waktu"]
    headers.extend(config.REGISTER_MAP.keys())
    headers.append("Status_Sync")
    return headers


def apply_header_style(ws, headers):
    """Menerapkan styling profesional pada baris header Excel."""
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align


def apply_status_style(cell, status):
    """Menerapkan warna khusus untuk status sinkronisasi."""
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if status == "SYNCED":
        cell.font = Font(name="Calibri", size=10, bold=True, color="006100")
        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    elif status == "PENDING":
        cell.font = Font(name="Calibri", size=10, bold=True, color="9C6500")
        cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def auto_fit_columns(ws):
    """Menyesuaikan lebar kolom agar tidak terpotong."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)


def reset_excel():
    """
    Membersihkan dan membuat ulang file Excel dengan header yang sesuai
    dengan REGISTER_MAP saat ini. Data lama akan dikosongkan agar bersih.

    Returns:
        bool: True jika berhasil.
    """
    excel_path = get_excel_path()
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PowerMeter_Data"

        headers = get_headers()
        apply_header_style(ws, headers)
        ws.freeze_panes = "A2"
        auto_fit_columns(ws)
        wb.save(excel_path)
        print(f"[INFO] File Excel berhasil dibersihkan & di-reset: {excel_path}")

        # Hapus juga file CSV lama jika ada agar tidak rancu
        csv_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "powermeter_data.csv")
        if os.path.exists(csv_path):
            try:
                os.remove(csv_path)
            except Exception:
                pass

        return True
    except PermissionError:
        print(f"[ERROR] Gagal mereset file: File Excel sedang dibuka di Microsoft Excel.")
        print(f"        Silakan tutup file {config.EXCEL_FILENAME} terlebih dahulu.")
        return False
    except Exception as e:
        print(f"[ERROR] Gagal mereset file Excel: {e}")
        return False


def ensure_excel_exists():
    """
    Memastikan file Excel ada dan headernya sesuai dengan konfigurasi terkini.
    Jika belum ada atau header berbeda dari REGISTER_MAP, buat ulang dengan rapi.
    Jika hanya kolom Status_Sync yang kurang, upgrade otomatis tanpa mereset data.

    Returns:
        str: Path file Excel.
    """
    excel_path = get_excel_path()
    expected_headers = get_headers()

    if not os.path.exists(excel_path):
        reset_excel()
        return excel_path

    # Jika file sudah ada, periksa apakah headernya masih cocok
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        existing_headers = [cell.value for cell in ws[1]]

        if existing_headers != expected_headers:
            # Jika hanya kurang kolom Status_Sync, upgrade tanpa mereset data
            headers_without_sync = [h for h in expected_headers if h != "Status_Sync"]
            if existing_headers == headers_without_sync:
                sync_col_idx = len(expected_headers)
                ws.cell(row=1, column=sync_col_idx, value="Status_Sync")
                apply_header_style(ws, expected_headers)

                for r in range(2, ws.max_row + 1):
                    c = ws.cell(row=r, column=sync_col_idx, value="SYNCED")
                    apply_status_style(c, "SYNCED")

                auto_fit_columns(ws)
                wb.save(excel_path)
                print("[INFO] Kolom 'Status_Sync' berhasil ditambahkan ke file Excel yang ada.")
            else:
                wb.close()
                print("[INFO] Header konfigurasi REGISTER_MAP telah berubah. Memperbarui struktur file Excel...")
                reset_excel()
                return excel_path
        wb.close()
    except Exception:
        pass

    return excel_path


# Alias untuk kompatibilitas
def ensure_csv_exists():
    return ensure_excel_exists()


def save_to_excel(data, sync_status="PENDING", recorded_at=None):
    """
    Menyimpan satu baris data pembacaan ke file Excel (.xlsx).
    Data dipetakan langsung berdasarkan nama kolom header agar nilai tidak pernah tertukar.

    Args:
        data: dict berisi {nama_parameter: nilai}
        sync_status (str): "SYNCED" atau "PENDING"
        recorded_at (datetime/str, optional): Waktu pencatatan data

    Returns:
        int: Nomor baris yang baru disimpan (truthy), atau 0 jika gagal.
    """
    excel_path = ensure_excel_exists()

    now = datetime.now() if recorded_at is None else (
        recorded_at if isinstance(recorded_at, datetime) else datetime.fromisoformat(str(recorded_at))
    )
    timestamp = now.strftime("%Y-%m-%d %H:%M:%S")
    tanggal = now.strftime("%Y-%m-%d")
    waktu = now.strftime("%H:%M:%S")

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        # Ambil header aktual dari baris 1
        headers = [cell.value for cell in ws[1]]

        # Susun baris sesuai urutan header persis
        row = []
        for h in headers:
            if h == "Timestamp":
                row.append(timestamp)
            elif h == "Tanggal":
                row.append(tanggal)
            elif h == "Waktu":
                row.append(waktu)
            elif h == "Status_Sync":
                row.append(sync_status)
            else:
                val = data.get(h)
                if val is not None:
                    try:
                        row.append(float(val))
                    except (ValueError, TypeError):
                        row.append(val)
                else:
                    row.append(0.0)

        # Append row ke sheet
        ws.append(row)
        current_row = ws.max_row

        center_align = Alignment(horizontal="center")
        right_align = Alignment(horizontal="right")

        for col_idx, (h_name, val) in enumerate(zip(headers, row), 1):
            cell = ws.cell(row=current_row, column=col_idx)
            if h_name in ("Timestamp", "Tanggal", "Waktu"):
                cell.alignment = center_align
            elif h_name == "Status_Sync":
                apply_status_style(cell, val)
            else:
                cell.alignment = right_align
                if isinstance(val, (int, float)):
                    cell.number_format = "0.000"

        wb.save(excel_path)
        return current_row

    except PermissionError:
        print(f"\n  [PERINGATAN] File '{config.EXCEL_FILENAME}' sedang dibuka di Microsoft Excel!")
        print("               Harap tutup file tersebut agar data baru bisa disimpan otomatis.\n")
        return 0
    except Exception as e:
        print(f"\n  [ERROR] Gagal menyimpan ke file Excel: {e}\n")
        return 0


# Alias untuk kompatibilitas fungsi lama
def save_to_csv(data):
    return save_to_excel(data) > 0


def update_row_sync_status(row_index, status="SYNCED"):
    """
    Memperbarui nilai kolom Status_Sync untuk baris tertentu.

    Args:
        row_index (int): Nomor baris di sheet (>= 2)
        status (str): "SYNCED" atau "PENDING"

    Returns:
        bool: True jika berhasil diperbarui.
    """
    if row_index < 2:
        return False

    excel_path = get_excel_path()
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        if "Status_Sync" not in headers:
            wb.close()
            return False

        col_idx = headers.index("Status_Sync") + 1
        cell = ws.cell(row=row_index, column=col_idx, value=status)
        apply_status_style(cell, status)

        wb.save(excel_path)
        return True
    except Exception as e:
        print(f"  [ERROR] Gagal memperbarui status baris #{row_index}: {e}")
        return False


def mark_rows_synced(row_indices):
    """
    Memperbarui sekumpulan nomor baris menjadi 'SYNCED' dalam satu kali save.

    Args:
        row_indices (list[int]): Daftar nomor baris

    Returns:
        int: Jumlah baris yang berhasil diupdate
    """
    if not row_indices:
        return 0

    excel_path = get_excel_path()
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        if "Status_Sync" not in headers:
            wb.close()
            return 0

        col_idx = headers.index("Status_Sync") + 1
        updated_count = 0

        for r in row_indices:
            if 2 <= r <= ws.max_row:
                cell = ws.cell(row=r, column=col_idx, value="SYNCED")
                apply_status_style(cell, "SYNCED")
                updated_count += 1

        wb.save(excel_path)
        return updated_count
    except Exception as e:
        print(f"  [ERROR] Gagal memperbarui status baris synced: {e}")
        return 0


def get_pending_rows():
    """
    Mengambil semua baris data dari file Excel yang memiliki Status_Sync == 'PENDING'.

    Returns:
        list[dict]: Daftar baris pending dengan format:
            [{
                "row_index": int,
                "timestamp": str,
                "data": dict parameter telemetri
            }, ...]
    """
    excel_path = get_excel_path()
    if not os.path.exists(excel_path):
        return []

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        if "Status_Sync" not in headers or "Timestamp" not in headers:
            wb.close()
            return []

        sync_col_idx = headers.index("Status_Sync")
        ts_col_idx = headers.index("Timestamp")

        pending_items = []
        for row_num, row_cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row_cells or all(c is None for c in row_cells):
                continue

            status_val = str(row_cells[sync_col_idx] or "").strip().upper()
            if status_val == "PENDING":
                ts_val = str(row_cells[ts_col_idx] or "")
                row_data = {}
                for h_name, cell_val in zip(headers, row_cells):
                    if h_name not in ("Timestamp", "Tanggal", "Waktu", "Status_Sync") and cell_val is not None:
                        try:
                            row_data[h_name] = float(cell_val)
                        except (ValueError, TypeError):
                            row_data[h_name] = cell_val

                pending_items.append({
                    "row_index": row_num,
                    "timestamp": ts_val,
                    "data": row_data,
                })

        wb.close()
        return pending_items
    except Exception as e:
        print(f"  [ERROR] Gagal membaca baris pending di Excel: {e}")
        return []


def get_pending_count():
    """Menghitung jumlah data pending yang belum tersinkronisasi."""
    return len(get_pending_rows())


def get_row_count():
    """
    Menghitung jumlah baris data (tidak termasuk header) dalam file Excel.

    Returns:
        int: Jumlah baris data, atau 0 jika file belum ada.
    """
    excel_path = get_excel_path()

    if not os.path.exists(excel_path):
        return 0

    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        ws = wb.active
        total_rows = 0
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                total_rows += 1
        wb.close()
        return max(0, total_rows - 1)  # Kurangi 1 untuk header
    except Exception:
        return 0


def get_last_n_rows(n=10):
    """
    Mengambil N baris terakhir dari file Excel.

    Args:
        n: Jumlah baris terakhir yang diambil.

    Returns:
        tuple: (headers, rows) — list header dan list of list data.
    """
    excel_path = get_excel_path()

    if not os.path.exists(excel_path):
        return [], []

    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        ws = wb.active

        all_rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                all_rows.append([str(c) if c is not None else "" for c in row])
        wb.close()

        if not all_rows:
            return [], []

        headers = all_rows[0]
        data_rows = all_rows[1:]
        return headers, data_rows[-n:]
    except Exception:
        return [], []


def get_file_info():
    """
    Mendapatkan informasi tentang file Excel termasuk statistik sync.

    Returns:
        dict: Informasi file Excel.
    """
    excel_path = get_excel_path()
    info = {
        "path": excel_path,
        "exists": os.path.exists(excel_path),
        "row_count": 0,
        "pending_count": 0,
        "synced_count": 0,
        "file_size": "0 B",
        "first_record": "-",
        "last_record": "-",
    }

    if not info["exists"]:
        return info

    try:
        size_bytes = os.path.getsize(excel_path)
        if size_bytes < 1024:
            info["file_size"] = f"{size_bytes} B"
        elif size_bytes < 1024 * 1024:
            info["file_size"] = f"{size_bytes / 1024:.1f} KB"
        else:
            info["file_size"] = f"{size_bytes / (1024 * 1024):.2f} MB"

        wb = openpyxl.load_workbook(excel_path, read_only=True)
        ws = wb.active
        data_rows = []
        headers = []

        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                if not headers:
                    headers = list(row)
                else:
                    data_rows.append(row)
        wb.close()

        if data_rows:
            info["row_count"] = len(data_rows)
            first_val = data_rows[0][0]
            last_val = data_rows[-1][0]
            info["first_record"] = str(first_val) if first_val is not None else "-"
            info["last_record"] = str(last_val) if last_val is not None else "-"

            if "Status_Sync" in headers:
                sync_idx = headers.index("Status_Sync")
                pending = sum(1 for r in data_rows if len(r) > sync_idx and str(r[sync_idx] or "").strip().upper() == "PENDING")
                info["pending_count"] = pending
                info["synced_count"] = len(data_rows) - pending

    except Exception:
        pass

    return info


# Alias untuk kompatibilitas
def get_csv_file_info():
    return get_file_info()
